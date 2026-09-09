
import openpyxl,json,subprocess,os
J=json.load(open('/tmp/claude-0/-home-user/e4610181-73c9-5b97-9776-a971b4d99e53/scratchpad/model.json')); XL='/home/user/Claude-septiembre/nocta/finanzas/modelo_bootstrap_nocta.xlsx'
res={}; os.makedirs('/tmp/lo/out',exist_ok=True)
for scen,name in [(1,'Pesimista'),(2,'Base'),(3,'Optimista')]:
    wb=openpyxl.load_workbook(XL); wb['Supuestos']['B2']=scen; tmp=f'/tmp/lo/in{scen}.xlsx'; wb.save(tmp)
    subprocess.run(['soffice','--headless','--calc','--convert-to','xlsx','--outdir','/tmp/lo/out',tmp],capture_output=True,timeout=180)
    w2=openpyxl.load_workbook(f'/tmp/lo/out/in{scen}.xlsx',data_only=True); M=w2['Modelo']; H=[c.value for c in M[1]]; idx={h:i for i,h in enumerate(H)}
    bad=[c.coordinate for r in M.iter_rows(min_row=2) for c in r if isinstance(c.value,str) and c.value.startswith('#')]
    mx=0
    for m in range(13):
        row=[c.value for c in M[m+2]]
        for xl,pk in [("Anuncios","spend"),("Pedidos","orders"),("Cobrado","revenue"),("Resultado","result"),("Caja fin","cash"),("Stock nariz fin","stock_n"),("IVA trimestre","vat_q"),("Aportación","aport")]:
            v=row[idx[xl]]; p=J['scen'][name][m][pk]
            if m==0 and xl=="Aportación": continue
            mx=max(mx,abs(float(v)-float(p)))
    res[name]=(mx,len(bad))
print('LibreOffice vs Python (max diff, error cells):',res)
AOV=J['AOV']; LAND=J['LAND']; checks=[]
for name,rows in J['scen'].items():
    e1=max(abs(r['orders']*AOV-r['revenue']) for r in rows)
    e2=max(abs(r['cpa']-r['cpc']/(LAND*r['cvr'])) for r in rows if r['cpa'])
    s=400+sum(r['aport'] for r in rows)+sum(r['revenue']-r['pay']-r['ret'] for r in rows)-sum(r['spend'] for r in rows)-sum(r['fixed'] for r in rows)-sum(r['reorder']*1.21 for r in rows)-sum(r['ship']*1.21 for r in rows)-sum(r['vat_q'] for r in rows)
    e3=abs(s-rows[-1]['cash']); minstock=min(min(r['stock_n'],r['stock_s']) for r in rows)
    e4=abs(125*2.90+40*1.60+sum(r['reorder'] for r in rows)-sum(r['cogs'] for r in rows)-rows[-1]['stock_val'])
    e5=abs(sum(r['vat']-r['iva_soportado'] for r in rows[:12])-sum(r['vat_q'] for r in rows))
    e6=max(abs(r['result']-(r['contrib']-r['spend']-r['fixed'])) for r in rows)
    checks.append((name,e1,e2,e3,minstock,e4,e5,e6))
for c in checks: print(c)
json.dump(dict(lo={k:v[0] for k,v in res.items()},errcells={k:v[1] for k,v in res.items()},checks=checks),open('/tmp/claude-0/-home-user/e4610181-73c9-5b97-9776-a971b4d99e53/scratchpad/checks.json','w'))
